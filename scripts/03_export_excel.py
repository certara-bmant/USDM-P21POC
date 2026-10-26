"""Export the generated P21 artefacts to a P21-style Excel spec for review."""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
P21, OUTX = sys.argv[1], sys.argv[2]
L = lambda f: json.load(open(os.path.join(P21, f)))
define, td, crf, so, cov = L('p21_sdtm_study_define.json'), L('p21_trial_design_domains.json'), L('p21_crf_form_spec.json'), L('p21_study_object.json'), L('coverage_report.json')
linked = L([f for f in os.listdir(P21) if f.endswith('_P21-linked.json')][0])
F = 'Arial'; HF = PatternFill('solid', fgColor='1F3864'); HFont = Font(name=F, bold=True, color='FFFFFF', size=10); BF = Font(name=F, size=10)
wb = Workbook(); wb.remove(wb.active)
def sheet(name, cols, rows, widths=None):
    ws = wb.create_sheet(name[:31]); ws.append(cols)
    for c in ws[1]: c.fill, c.font, c.alignment = HF, HFont, Alignment(vertical='center', wrap_text=True)
    for r in rows:
        ws.append([(json.dumps(v) if isinstance(v, (dict, list)) else v) for v in r])
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = BF
    for i, col in enumerate(cols, 1):
        longest = max([len(str(r[i - 1] or '')) for r in rows[:200]] or [0])
        w = (widths or {}).get(col) or min(60, max(10, len(str(col)) + 2, longest + 2))
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    return ws
ca = lambda o, k: (o.get('customAttribute') or {}).get(k)

# Summary with live COUNTA formulas over the detail sheets
ws = wb.create_sheet('Summary')
ws['A1'] = 'CTX101-001 — USDM -> P21 generation POC'; ws['A1'].font = Font(name=F, bold=True, size=14)
ws['A2'] = 'Fictional Phase 1 SAD study. All counts below are formulas over the detail tabs.'; ws['A2'].font = Font(name=F, italic=True, size=10)
items = [('SDTM datasets', 'Datasets'), ('SDTM variables', 'Variables'), ('Value-level rows', 'ValueLevel'), ('Study codelists', 'Codelists'), ('Study terms', 'Terms'),
         ('TA rows', 'TA'), ('TE rows', 'TE'), ('TV rows', 'TV'), ('TI rows', 'TI'), ('TS rows', 'TS'),
         ('CRF events', 'CRF_Events'), ('CRF forms', 'CRF_Forms'), ('CRF SoA rows', 'CRF_SoA'), ('CRF questions', 'CRF_Questions'), ('CRF codelists', 'CRF_Codelists'),
         ('BCs in USDM', 'BC_Linkage'), ('Coverage findings', 'Coverage')]
ws['A4'], ws['B4'] = 'Artefact', 'Count'
for c in ws[4]: c.fill, c.font = HF, HFont
for i, (lab, sh) in enumerate(items, 5):
    ws.cell(i, 1, lab).font = BF; ws.cell(i, 2, f"=COUNTA('{sh}'!{'B:B' if sh in ('Variables', 'ValueLevel', 'Terms') else 'A:A'})-1").font = BF
r0 = 5 + len(items) + 1
ws.cell(r0, 1, 'Generation options').font = Font(name=F, bold=True)
for j, (k, v) in enumerate(define['standard']['customAttribute']['options'].items(), 1):
    ws.cell(r0 + j, 1, k).font = BF; ws.cell(r0 + j, 2, str(v)).font = BF
ws.cell(r0 + 4, 1, 'Master standard / terminology').font = Font(name=F, bold=True)
ws.cell(r0 + 5, 1, so['dataPackages'][1]['Standard']).font = BF; ws.cell(r0 + 6, 1, so['dataPackages'][1]['ControlledTerminology']).font = BF
ws.column_dimensions['A'].width = 48; ws.column_dimensions['B'].width = 14

s = so['studyObject']
sheet('Study', ['Field', 'Value'], [[k, ', '.join(v) if isinstance(v, list) else v] for k, v in s.items() if k != 'customAttribute'] +
      [[f"DataPackage: {p['Name']}", f"Standard={p['Standard']}; CT={p['ControlledTerminology']}"] for p in so['dataPackages']], {'Field': 26, 'Value': 110})
sheet('Datasets', ['Dataset', 'Label', 'Class', 'Structure', 'Key Variables', 'Repeating', 'Reference Data', 'Source internalId', 'Specializations'],
      [[d['dataset'], d['label'], d['class'], d['structure'], d['keyVariables'], d['repeating'], d['referenceData'], ca(d, 'sourceInternalId'), ', '.join(ca(d, 'specializations') or [])] for d in define['datasets']])
sheet('Variables', ['Order', 'Dataset', 'Variable', 'Label', 'Data Type', 'Length', 'Core', 'Codelist', 'Origin', 'Role', 'Inclusion', 'Source internalId', 'Specializations'],
      [[(int(v['order']) if v['order'] else None), v['dataset'], v['variable'], v['label'], v['dataType'], v['length'], v['core'], v['codelist'], v['origin'], v['role'],
        ca(v, 'inclusion'), ca(v, 'sourceInternalId'), ', '.join(ca(v, 'specializations') or [])] for v in define['variables']])
sheet('ValueLevel', ['Order', 'Dataset', 'Variable', 'Where Clause', 'Label', 'Data Type', 'Length', 'Significant Digits', 'Mandatory', 'Assigned Value', 'Codelist', 'Origin', 'Source', 'BC', 'Specialization'],
      [[(int(v['order']) if v['order'] else None), v['dataset'], v['variable'], v['whereClause'], v['label'], v['dataType'], v['length'], v['significantDigits'], v['mandatory'], v['assignedValue'], v['codelist'],
        v['origin'], v['source'], ca(v, 'bcId'), ca(v, 'specialization')] for v in define['value_level']], {'Where Clause': 60})
sheet('Codelists', ['ID', 'Name', 'NCI Code', 'Data Type', 'Extensible', 'Parent Codelist', 'Subset Rule', 'Source internalId', 'Used By'],
      [[c['id'], c['name'], c['nciCode'], c['dataType'], c['extensible'], ca(c, 'parentCodelist'), ca(c, 'subsetRule'), ca(c, 'sourceInternalId'), ', '.join(ca(c, 'usedBy') or [])] for c in define['codelists']])
sheet('Terms', ['Codelist', 'Order', 'Term', 'Decoded Value', 'NCI Code', 'Sponsor Extended', 'Source internalId'],
      [[t['codelist'], int(t['order']), t['term'], t['decodedValue'], t['nciCode'], 'Yes' if ca(t, 'extended') else 'No', ca(t, 'sourceInternalId')] for t in define['terms']])
for dom, rows in td.items():
    cols = [k for k in rows[0] if k != '_usdm'] + ['USDM source'] if rows else ['STUDYID']
    sheet(dom, cols, [[r[k] for k in cols[:-1]] + [r.get('_usdm')] for r in rows])
def dump(name, rows):
    cols = list(rows[0].keys()) if rows else ['empty']
    sheet(name, cols, [[r.get(k) for k in cols] for r in rows])
dump('CRF_Events', crf['events']); dump('CRF_Forms', crf['forms']); dump('CRF_SoA', crf['scheduleOfActivities'])
dump('CRF_Questions', crf['questions']); dump('CRF_Codelists', crf['codelists']); dump('CRF_Terms', crf['terms'])
# BC linkage: the hidden persistence layer
SD, CR, P = 'http://www.cdisc.org/usdm/extensions/specializations/sdtm', 'http://www.cdisc.org/usdm/extensions/specializations/crf', 'https://www.certara.com/usdm/extensions/p21'
sv = linked['study']['versions'][0]; acts = {}
for a in sv['studyDesigns'][0]['activities']:
    for b in a['biomedicalConceptIds']: acts.setdefault(b, []).append(a['name'])
rows = []
for b in sv['biomedicalConcepts']:
    e = {x['url']: x for x in b['extensionAttributes']}
    link = {x['url'].split('/')[-1]: x.get('valueString') for x in e.get(f'{P}/sdtm-link', {}).get('extensionAttributes', [])}
    rows.append([b['id'], b['name'], b['code']['standardCode']['code'], ', '.join(acts.get(b['id'], [])), (e.get(SD) or {}).get('valueString'), (e.get(CR) or {}).get('valueString'),
                 link.get('dataset'), link.get('topicVariable'), link.get('topicTerm'), link.get('whereClause'), len(b['properties'])])
sheet('BC_Linkage', ['USDM BC id', 'BC name', 'BC C-code', 'Activities', 'SDTM specialization ext', 'CRF specialization ext', 'P21 dataset (internalId|key)',
                     'P21 topic variable', 'P21 topic term', 'Where clause', '# properties'], rows)
cr = []
for k, v in cov.items():
    if k == 'summary': continue
    for x in v: cr.append([k, json.dumps(x) if isinstance(x, dict) else x])
sheet('Coverage', ['Finding type', 'Detail'], cr, {'Finding type': 48, 'Detail': 140})
wb.move_sheet('Summary', -(len(wb.sheetnames) - 1))
wb.save(OUTX); print('saved', OUTX, wb.sheetnames)
